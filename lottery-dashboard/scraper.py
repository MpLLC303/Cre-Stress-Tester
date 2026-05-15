"""Colorado Lottery scratch ticket dashboard data pipeline.

Fetches the official Scratch Insider XLSX feed, enriches each active game
with its per-tier prize breakdown from the game page, computes derived
columns (including an EV score), and renders index.html from the Jinja
template. If the live endpoints are unreachable, a sample dataset is used
so the dashboard still renders for visual verification.
"""

from __future__ import annotations

import datetime as dt
import email.utils
import json
import logging
import re
import sys
import time
import traceback
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd
import requests
from bs4 import BeautifulSoup
from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
GAME_PAGES = DATA / "game_pages"
LOGS = ROOT / "logs"
TEMPLATES = ROOT / "templates"
OUTPUT_HTML = ROOT / "index.html"
STATUS_HTML = ROOT / "status.html"
CURRENT_JSON = DATA / "current.json"
HISTORY_JSONL = DATA / "history.jsonl"
REFRESH_LOG = DATA / "refresh_log.jsonl"

XLSX_URL = "https://www.coloradolottery.com/en/player-tools/scratch-insider/?xlsx="
INSIDER_HTML_URL = "https://www.coloradolottery.com/en/player-tools/scratch-insider/"
GAME_URL_TEMPLATE = "https://www.coloradolottery.com/en/games/scratch/game/{slug}/"
WINNING_STORES_TEMPLATE = (
    "https://www.coloradolottery.com/en/player-tools/winning-stores/"
    "?zip=80501&radius=5&game_type=scratch&game_number={game_number}"
)

USER_AGENT = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"
)
REQUEST_TIMEOUT = 30
POLITE_DELAY_SECONDS = 1.0


for d in (DATA, GAME_PAGES, LOGS):
    d.mkdir(parents=True, exist_ok=True)

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(message)s",
    handlers=[
        logging.FileHandler(LOGS / "scraper.log"),
        logging.StreamHandler(sys.stdout),
    ],
)
log = logging.getLogger("scraper")


@dataclass
class FetchResult:
    ok: bool
    source: str
    note: str = ""


def http_session() -> requests.Session:
    s = requests.Session()
    s.headers.update({"User-Agent": USER_AGENT, "Accept": "*/*"})
    return s


def slugify(name: str) -> str:
    s = name.lower().strip()
    s = s.replace("&", "and")
    s = re.sub(r"[‘’“”]", "", s)
    s = re.sub(r"[^a-z0-9]+", "-", s)
    return s.strip("-")


def fetch_xlsx(session: requests.Session, dest: Path) -> FetchResult:
    log.info("Fetching XLSX feed: %s", XLSX_URL)
    try:
        r = session.get(XLSX_URL, timeout=REQUEST_TIMEOUT)
    except requests.RequestException as exc:
        return FetchResult(False, "xlsx", f"request error: {exc}")
    if r.status_code != 200:
        return FetchResult(False, "xlsx", f"http {r.status_code}")
    ctype = r.headers.get("Content-Type", "")
    body = r.content
    if not (body[:2] == b"PK" or "spreadsheet" in ctype or "octet-stream" in ctype):
        return FetchResult(False, "xlsx", f"non-binary response ({ctype})")
    dest.write_bytes(body)
    log.info("Saved XLSX (%d bytes) to %s", len(body), dest)
    return FetchResult(True, "xlsx")


def parse_xlsx(path: Path) -> pd.DataFrame:
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return pd.DataFrame()
    header_idx = 0
    for i, row in enumerate(rows[:10]):
        joined = " ".join(str(c) for c in row if c is not None).lower()
        if "game" in joined and ("number" in joined or "#" in joined):
            header_idx = i
            break
    raw_headers = [str(c).strip() if c is not None else "" for c in rows[header_idx]]
    data_rows = rows[header_idx + 1 :]
    df = pd.DataFrame(data_rows, columns=raw_headers)
    df = df.dropna(how="all")
    df.columns = [_normalize_header(c) for c in df.columns]
    return _coerce_columns(df)


def _normalize_header(h: str) -> str:
    h = h.lower().strip()
    mapping = {
        "game name": "game_name",
        "game": "game_name",
        "game number": "game_number",
        "game #": "game_number",
        "ticket price": "ticket_price",
        "price": "ticket_price",
        "game start": "game_start",
        "start date": "game_start",
        "last day to claim": "last_day_to_claim",
        "top prize": "top_prize",
        "total top prizes": "total_top_prizes",
        "top prizes remaining": "top_prizes_remaining",
        "overall odds": "overall_odds",
        "number of eligible drawings": "num_eligible_drawings",
        "num eligible drawings": "num_eligible_drawings",
        "payout percentage": "payout_percentage",
        "payout %": "payout_percentage",
    }
    return mapping.get(h, re.sub(r"[^a-z0-9]+", "_", h).strip("_"))


def _to_money(val: Any) -> float | None:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).replace("$", "").replace(",", "").strip()
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _to_int(val: Any) -> int | None:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return None
    if isinstance(val, (int, float)):
        return int(val)
    s = str(val).replace(",", "").strip()
    if not s or s.lower() in {"not set", "n/a", "-"}:
        return None
    try:
        return int(float(s))
    except ValueError:
        return None


def _to_date(val: Any) -> str | None:
    if val is None:
        return None
    if isinstance(val, (dt.datetime, dt.date)):
        return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    if not s or s.lower() == "not set":
        return None
    for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%m-%d-%Y", "%b %d, %Y"):
        try:
            return dt.datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def _to_percent(val: Any) -> float | None:
    if val is None:
        return None
    if isinstance(val, (int, float)) and not pd.isna(val):
        v = float(val)
        return v if v <= 1.5 else v / 100.0
    s = str(val).strip().rstrip("%")
    try:
        v = float(s)
        return v / 100.0 if v > 1.5 else v
    except ValueError:
        return None


def _coerce_columns(df: pd.DataFrame) -> pd.DataFrame:
    expected = [
        "game_name",
        "game_number",
        "ticket_price",
        "game_start",
        "last_day_to_claim",
        "top_prize",
        "total_top_prizes",
        "top_prizes_remaining",
        "overall_odds",
        "num_eligible_drawings",
        "payout_percentage",
    ]
    for col in expected:
        if col not in df.columns:
            df[col] = None
    df["game_name"] = df["game_name"].astype(str).str.strip()
    df["game_number"] = df["game_number"].apply(_to_int)
    df["ticket_price"] = df["ticket_price"].apply(_to_money)
    df["game_start"] = df["game_start"].apply(_to_date)
    df["last_day_to_claim"] = df["last_day_to_claim"].apply(_to_date)
    df["top_prize"] = df["top_prize"].apply(_to_money)
    df["total_top_prizes"] = df["total_top_prizes"].apply(_to_int)
    df["top_prizes_remaining"] = df["top_prizes_remaining"].apply(_to_int)
    df["num_eligible_drawings"] = df["num_eligible_drawings"].apply(_to_int)
    df["payout_percentage"] = df["payout_percentage"].apply(_to_percent)
    df["overall_odds"] = df["overall_odds"].astype(str).str.strip()
    df = df[df["game_number"].notna()].reset_index(drop=True)
    return df[expected]


def parse_insider_html(html: str) -> pd.DataFrame:
    soup = BeautifulSoup(html, "html.parser")
    table = soup.find("table")
    if not table:
        return pd.DataFrame()
    headers = [th.get_text(strip=True) for th in table.find_all("th")]
    rows = []
    for tr in table.find_all("tr"):
        cells = [td.get_text(" ", strip=True) for td in tr.find_all("td")]
        if cells:
            rows.append(cells)
    if not rows:
        return pd.DataFrame()
    df = pd.DataFrame(rows, columns=headers[: len(rows[0])])
    df.columns = [_normalize_header(c) for c in df.columns]
    return _coerce_columns(df)


def fetch_insider_html(session: requests.Session) -> FetchResult:
    log.info("Falling back to insider HTML page")
    try:
        r = session.get(INSIDER_HTML_URL, timeout=REQUEST_TIMEOUT)
    except requests.RequestException as exc:
        return FetchResult(False, "html", f"request error: {exc}")
    if r.status_code != 200:
        return FetchResult(False, "html", f"http {r.status_code}")
    (DATA / "insider_fallback.html").write_text(r.text, encoding="utf-8")
    return FetchResult(True, "html")


def fetch_game_page(session: requests.Session, game_number: int, slug: str) -> str | None:
    cached = GAME_PAGES / f"{game_number}.html"
    meta = GAME_PAGES / f"{game_number}.meta.json"
    headers = {}
    if meta.exists():
        try:
            saved = json.loads(meta.read_text())
            if saved.get("last_modified"):
                headers["If-Modified-Since"] = saved["last_modified"]
            if saved.get("etag"):
                headers["If-None-Match"] = saved["etag"]
        except json.JSONDecodeError:
            pass
    url = GAME_URL_TEMPLATE.format(slug=slug)
    try:
        r = session.get(url, headers=headers, timeout=REQUEST_TIMEOUT)
    except requests.RequestException as exc:
        log.warning("game page fetch failed for %s: %s", slug, exc)
        return cached.read_text(encoding="utf-8") if cached.exists() else None
    time.sleep(POLITE_DELAY_SECONDS)
    if r.status_code == 304 and cached.exists():
        log.info("game %s unchanged (304)", game_number)
        return cached.read_text(encoding="utf-8")
    if r.status_code != 200:
        log.warning("game %s http %s", game_number, r.status_code)
        return cached.read_text(encoding="utf-8") if cached.exists() else None
    cached.write_text(r.text, encoding="utf-8")
    meta.write_text(
        json.dumps(
            {
                "last_modified": r.headers.get("Last-Modified", ""),
                "etag": r.headers.get("ETag", ""),
                "fetched_at": dt.datetime.utcnow().isoformat() + "Z",
            }
        )
    )
    return r.text


def parse_prize_tiers(html: str) -> list[dict[str, Any]]:
    soup = BeautifulSoup(html, "html.parser")
    tiers: list[dict[str, Any]] = []
    for table in soup.find_all("table"):
        head = " ".join(th.get_text(" ", strip=True).lower() for th in table.find_all("th"))
        if "prize" not in head:
            continue
        if "remaining" not in head and "left" not in head:
            continue
        for tr in table.find_all("tr"):
            cells = [td.get_text(" ", strip=True) for td in tr.find_all("td")]
            if len(cells) < 3:
                continue
            prize = _to_money(cells[0])
            total = _to_int(cells[1])
            remaining = _to_int(cells[2])
            if prize is None and total is None and remaining is None:
                continue
            tiers.append({"prize": prize, "total": total, "remaining": remaining})
        if tiers:
            break
    return tiers


def compute_derived(df: pd.DataFrame, tier_map: dict[int, list[dict[str, Any]]]) -> pd.DataFrame:
    today = dt.date.today()
    df = df.copy()
    derived = {
        "days_since_launch": [],
        "pct_top_prizes_remaining": [],
        "pct_total_prize_pool_remaining": [],
        "days_to_claim_deadline": [],
        "conditional_ev_per_dollar": [],
        "ev_score": [],
        "prize_tiers": [],
        "game_slug": [],
        "winning_stores_url": [],
    }
    for _, row in df.iterrows():
        start = _parse_iso(row["game_start"])
        days_since = (today - start).days if start else None
        derived["days_since_launch"].append(days_since)

        total_top = row["total_top_prizes"] or 0
        remaining_top = row["top_prizes_remaining"] or 0
        pct_top = (remaining_top / total_top) if total_top else None
        derived["pct_top_prizes_remaining"].append(pct_top)

        tiers = tier_map.get(int(row["game_number"]), [])
        if tiers:
            total_value = sum((t["total"] or 0) * (t["prize"] or 0) for t in tiers)
            remaining_value = sum((t["remaining"] or 0) * (t["prize"] or 0) for t in tiers)
            pct_pool = (remaining_value / total_value) if total_value else None
        else:
            pct_pool = pct_top
        derived["pct_total_prize_pool_remaining"].append(pct_pool)

        claim = _parse_iso(row["last_day_to_claim"])
        days_to_claim = (claim - today).days if claim else None
        derived["days_to_claim_deadline"].append(days_to_claim)

        payout = row["payout_percentage"]
        if payout is not None and pct_pool is not None:
            cond_ev = float(payout) * float(pct_pool)
        else:
            cond_ev = None
        derived["conditional_ev_per_dollar"].append(cond_ev)

        ev = cond_ev
        if ev is not None:
            if days_to_claim is not None:
                if days_to_claim < 90:
                    ev *= 0.75
                elif days_to_claim < 180:
                    ev *= 0.90
            if days_since is not None and days_since > 365:
                ev *= 0.95
        derived["ev_score"].append(ev)

        derived["prize_tiers"].append(tiers)
        slug = slugify(row["game_name"])
        derived["game_slug"].append(slug)
        derived["winning_stores_url"].append(
            WINNING_STORES_TEMPLATE.format(game_number=int(row["game_number"]))
        )

    for k, v in derived.items():
        df[k] = v
    return df


def _parse_iso(s: str | None) -> dt.date | None:
    if not s:
        return None
    try:
        return dt.datetime.strptime(s, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return None


def build_summary(df: pd.DataFrame) -> dict[str, Any]:
    by_denom: dict[str, dict[str, float]] = {}
    for price, group in df.dropna(subset=["ticket_price"]).groupby("ticket_price"):
        payouts = group["payout_percentage"].dropna()
        by_denom[f"${int(price)}"] = {
            "count": int(len(group)),
            "avg_payout": float(payouts.mean()) if not payouts.empty else None,
        }
    top3 = (
        df.dropna(subset=["ev_score"])
        .sort_values("ev_score", ascending=False)
        .head(3)[["game_name", "game_number", "ticket_price", "ev_score"]]
        .to_dict(orient="records")
    )
    total_remaining_pool = 0.0
    for tiers in df["prize_tiers"]:
        for t in tiers or []:
            total_remaining_pool += (t.get("remaining") or 0) * (t.get("prize") or 0)
    return {
        "by_denomination": by_denom,
        "top3_by_ev": top3,
        "total_remaining_prize_pool": total_remaining_pool,
        "active_games": int(len(df)),
    }


def append_history(df: pd.DataFrame, today: str) -> None:
    keep = [
        "game_number",
        "game_name",
        "ticket_price",
        "top_prizes_remaining",
        "total_top_prizes",
        "pct_total_prize_pool_remaining",
        "ev_score",
        "payout_percentage",
    ]
    with HISTORY_JSONL.open("a", encoding="utf-8") as f:
        for _, row in df.iterrows():
            snapshot = {"snapshot_date": today}
            for k in keep:
                v = row.get(k)
                if isinstance(v, float) and pd.isna(v):
                    v = None
                snapshot[k] = v
            f.write(json.dumps(snapshot, default=str) + "\n")


def render_html(df: pd.DataFrame, summary: dict[str, Any], source: str, fallback_note: str) -> None:
    env = Environment(
        loader=FileSystemLoader(str(TEMPLATES)),
        autoescape=select_autoescape(["html"]),
    )
    template = env.get_template("dashboard.html.j2")
    games = json.loads(df.to_json(orient="records"))
    now = dt.datetime.now()
    payload = {
        "generated_at": now.strftime("%Y-%m-%d %H:%M:%S %Z").strip(),
        "generated_iso": now.isoformat(timespec="seconds"),
        "games": games,
        "summary": summary,
        "source": source,
        "fallback_note": fallback_note,
        "refresh_log": _read_refresh_log(),
    }
    html = template.render(**payload, data_json=json.dumps({"games": games, "summary": summary}))
    OUTPUT_HTML.write_text(html, encoding="utf-8")
    CURRENT_JSON.write_text(json.dumps(payload, indent=2, default=str), encoding="utf-8")
    log.info("Wrote %s (%d games)", OUTPUT_HTML, len(games))


def _read_refresh_log() -> list[dict[str, Any]]:
    if not REFRESH_LOG.exists():
        return []
    rows = []
    for line in REFRESH_LOG.read_text(encoding="utf-8").splitlines()[-20:]:
        try:
            rows.append(json.loads(line))
        except json.JSONDecodeError:
            continue
    return list(reversed(rows))


def write_refresh_log(entry: dict[str, Any]) -> None:
    with REFRESH_LOG.open("a", encoding="utf-8") as f:
        f.write(json.dumps(entry, default=str) + "\n")


def write_status_page(error: str) -> None:
    last = "(none)"
    if CURRENT_JSON.exists():
        try:
            last = json.loads(CURRENT_JSON.read_text())["generated_at"]
        except Exception:
            pass
    STATUS_HTML.write_text(
        f"""<!doctype html>
<html><head><meta charset="utf-8"><title>Lottery dashboard status</title>
<style>body{{background:#07060A;color:#e6e3da;font-family:system-ui;padding:40px;}}
.banner{{background:#3a1a1a;border:1px solid #d4af37;padding:24px;border-radius:8px;}}</style>
</head><body><div class="banner">
<h1>Refresh failed</h1>
<p>Last successful refresh: <strong>{last}</strong></p>
<pre>{error}</pre>
</div></body></html>""",
        encoding="utf-8",
    )


SAMPLE_GAMES = [
    {
        "game_name": "Best Chance To Be A Millionaire",
        "game_number": 379,
        "ticket_price": 50.0,
        "game_start": "2025-08-12",
        "last_day_to_claim": None,
        "top_prize": 5000000.0,
        "total_top_prizes": 4,
        "top_prizes_remaining": 3,
        "overall_odds": "1 in 2.94",
        "num_eligible_drawings": 1,
        "payout_percentage": 0.78,
        "tiers": [
            {"prize": 5000000, "total": 4, "remaining": 3},
            {"prize": 100000, "total": 30, "remaining": 22},
            {"prize": 10000, "total": 200, "remaining": 140},
            {"prize": 1000, "total": 5000, "remaining": 3400},
            {"prize": 500, "total": 25000, "remaining": 16800},
            {"prize": 100, "total": 250000, "remaining": 170000},
            {"prize": 50, "total": 800000, "remaining": 540000},
        ],
    },
    {
        "game_name": "$3,000,000 Millionaire Maker",
        "game_number": 380,
        "ticket_price": 50.0,
        "game_start": "2024-11-04",
        "last_day_to_claim": "2026-08-15",
        "top_prize": 3000000.0,
        "total_top_prizes": 6,
        "top_prizes_remaining": 2,
        "overall_odds": "1 in 3.04",
        "num_eligible_drawings": 1,
        "payout_percentage": 0.77,
        "tiers": [
            {"prize": 3000000, "total": 6, "remaining": 2},
            {"prize": 100000, "total": 45, "remaining": 18},
            {"prize": 10000, "total": 220, "remaining": 95},
            {"prize": 1000, "total": 5500, "remaining": 2100},
            {"prize": 500, "total": 28000, "remaining": 9800},
            {"prize": 100, "total": 280000, "remaining": 95000},
            {"prize": 50, "total": 900000, "remaining": 310000},
        ],
    },
    {
        "game_name": "SET FOR LIFE",
        "game_number": 387,
        "ticket_price": 50.0,
        "game_start": "2026-02-10",
        "last_day_to_claim": None,
        "top_prize": 4000000.0,
        "total_top_prizes": 3,
        "top_prizes_remaining": 3,
        "overall_odds": "1 in 2.88",
        "num_eligible_drawings": 1,
        "payout_percentage": 0.79,
        "tiers": [
            {"prize": 4000000, "total": 3, "remaining": 3},
            {"prize": 100000, "total": 25, "remaining": 24},
            {"prize": 10000, "total": 180, "remaining": 170},
            {"prize": 1000, "total": 4800, "remaining": 4550},
            {"prize": 500, "total": 24000, "remaining": 22600},
            {"prize": 100, "total": 240000, "remaining": 225000},
            {"prize": 50, "total": 780000, "remaining": 730000},
        ],
    },
    {
        "game_name": "Diamond Dazzler",
        "game_number": 372,
        "ticket_price": 20.0,
        "game_start": "2024-05-20",
        "last_day_to_claim": "2026-09-30",
        "top_prize": 1000000.0,
        "total_top_prizes": 8,
        "top_prizes_remaining": 1,
        "overall_odds": "1 in 3.21",
        "num_eligible_drawings": 1,
        "payout_percentage": 0.72,
        "tiers": [
            {"prize": 1000000, "total": 8, "remaining": 1},
            {"prize": 50000, "total": 40, "remaining": 6},
            {"prize": 1000, "total": 4000, "remaining": 350},
            {"prize": 100, "total": 200000, "remaining": 22000},
            {"prize": 20, "total": 600000, "remaining": 72000},
        ],
    },
    {
        "game_name": "Holiday Cash Blowout",
        "game_number": 384,
        "ticket_price": 10.0,
        "game_start": "2025-11-01",
        "last_day_to_claim": "2026-07-15",
        "top_prize": 200000.0,
        "total_top_prizes": 12,
        "top_prizes_remaining": 9,
        "overall_odds": "1 in 3.48",
        "num_eligible_drawings": 1,
        "payout_percentage": 0.70,
        "tiers": [
            {"prize": 200000, "total": 12, "remaining": 9},
            {"prize": 10000, "total": 60, "remaining": 38},
            {"prize": 500, "total": 4500, "remaining": 2700},
            {"prize": 50, "total": 90000, "remaining": 54000},
            {"prize": 10, "total": 700000, "remaining": 420000},
        ],
    },
    {
        "game_name": "Cash Crossword",
        "game_number": 365,
        "ticket_price": 5.0,
        "game_start": "2023-09-15",
        "last_day_to_claim": "2026-06-01",
        "top_prize": 100000.0,
        "total_top_prizes": 10,
        "top_prizes_remaining": 1,
        "overall_odds": "1 in 3.95",
        "num_eligible_drawings": 1,
        "payout_percentage": 0.65,
        "tiers": [
            {"prize": 100000, "total": 10, "remaining": 1},
            {"prize": 1000, "total": 200, "remaining": 14},
            {"prize": 100, "total": 6000, "remaining": 480},
            {"prize": 25, "total": 40000, "remaining": 3200},
            {"prize": 5, "total": 500000, "remaining": 38000},
        ],
    },
    {
        "game_name": "Triple Lucky 7s",
        "game_number": 381,
        "ticket_price": 3.0,
        "game_start": "2025-06-03",
        "last_day_to_claim": None,
        "top_prize": 30000.0,
        "total_top_prizes": 20,
        "top_prizes_remaining": 14,
        "overall_odds": "1 in 4.22",
        "num_eligible_drawings": 1,
        "payout_percentage": 0.63,
        "tiers": [
            {"prize": 30000, "total": 20, "remaining": 14},
            {"prize": 500, "total": 800, "remaining": 520},
            {"prize": 50, "total": 12000, "remaining": 7800},
            {"prize": 6, "total": 300000, "remaining": 195000},
        ],
    },
    {
        "game_name": "$2 Bingo",
        "game_number": 376,
        "ticket_price": 2.0,
        "game_start": "2025-01-14",
        "last_day_to_claim": None,
        "top_prize": 20000.0,
        "total_top_prizes": 15,
        "top_prizes_remaining": 10,
        "overall_odds": "1 in 4.10",
        "num_eligible_drawings": 1,
        "payout_percentage": 0.62,
        "tiers": [
            {"prize": 20000, "total": 15, "remaining": 10},
            {"prize": 200, "total": 1200, "remaining": 720},
            {"prize": 20, "total": 18000, "remaining": 11000},
            {"prize": 4, "total": 240000, "remaining": 148000},
        ],
    },
    {
        "game_name": "Hit $500",
        "game_number": 369,
        "ticket_price": 1.0,
        "game_start": "2024-08-05",
        "last_day_to_claim": "2026-05-30",
        "top_prize": 500.0,
        "total_top_prizes": 800,
        "top_prizes_remaining": 120,
        "overall_odds": "1 in 4.55",
        "num_eligible_drawings": 1,
        "payout_percentage": 0.60,
        "tiers": [
            {"prize": 500, "total": 800, "remaining": 120},
            {"prize": 50, "total": 4000, "remaining": 700},
            {"prize": 5, "total": 60000, "remaining": 10000},
            {"prize": 1, "total": 600000, "remaining": 95000},
        ],
    },
]


def load_sample() -> tuple[pd.DataFrame, dict[int, list[dict[str, Any]]]]:
    rows = []
    tiers: dict[int, list[dict[str, Any]]] = {}
    for g in SAMPLE_GAMES:
        rows.append(
            {
                "game_name": g["game_name"],
                "game_number": g["game_number"],
                "ticket_price": g["ticket_price"],
                "game_start": g["game_start"],
                "last_day_to_claim": g["last_day_to_claim"],
                "top_prize": g["top_prize"],
                "total_top_prizes": g["total_top_prizes"],
                "top_prizes_remaining": g["top_prizes_remaining"],
                "overall_odds": g["overall_odds"],
                "num_eligible_drawings": g["num_eligible_drawings"],
                "payout_percentage": g["payout_percentage"],
            }
        )
        tiers[g["game_number"]] = g["tiers"]
    return pd.DataFrame(rows), tiers


def main() -> int:
    today = dt.date.today().isoformat()
    started = dt.datetime.utcnow().isoformat() + "Z"
    session = http_session()
    source = "unknown"
    fallback_note = ""
    df: pd.DataFrame
    tier_map: dict[int, list[dict[str, Any]]] = {}

    xlsx_path = DATA / f"scratch_insider_{today}.xlsx"
    result = fetch_xlsx(session, xlsx_path)
    if result.ok:
        try:
            df = parse_xlsx(xlsx_path)
            source = "xlsx"
        except Exception as exc:
            log.warning("xlsx parse failed: %s", exc)
            result = FetchResult(False, "xlsx", f"parse error: {exc}")

    if not result.ok:
        fallback_note = f"xlsx unavailable ({result.note}); trying HTML fallback"
        log.warning(fallback_note)
        html_result = fetch_insider_html(session)
        if html_result.ok:
            df = parse_insider_html((DATA / "insider_fallback.html").read_text(encoding="utf-8"))
            source = "html"
        else:
            note = f"both feeds failed (xlsx: {result.note}; html: {html_result.note})"
            log.error(note)
            if OUTPUT_HTML.exists():
                log.info("Leaving prior index.html intact; writing status.html")
                write_status_page(note)
                write_refresh_log({"started": started, "ok": False, "source": "none", "note": note})
                return 2
            log.info("No prior index.html; rendering with sample dataset for first-run visibility")
            df, tier_map = load_sample()
            source = "sample"
            fallback_note = note + "; rendered from bundled sample data"

    if source in ("xlsx", "html"):
        for _, row in df.iterrows():
            slug = slugify(row["game_name"])
            html = fetch_game_page(session, int(row["game_number"]), slug)
            if html:
                tier_map[int(row["game_number"])] = parse_prize_tiers(html)

    df = compute_derived(df, tier_map)
    df = df.sort_values("ev_score", ascending=False, na_position="last").reset_index(drop=True)

    summary = build_summary(df)
    append_history(df, today)
    render_html(df, summary, source, fallback_note)
    write_refresh_log(
        {
            "started": started,
            "finished": dt.datetime.utcnow().isoformat() + "Z",
            "ok": True,
            "source": source,
            "games": int(len(df)),
            "note": fallback_note,
        }
    )
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception:
        err = traceback.format_exc()
        log.error("scraper crashed: %s", err)
        write_status_page(err)
        write_refresh_log(
            {"started": dt.datetime.utcnow().isoformat() + "Z", "ok": False, "note": str(err)[:500]}
        )
        sys.exit(1)
